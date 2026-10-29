"""Excel review export with bounded sheet count and verified rich text."""
from collections import Counter
from copy import copy
import re

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from automatic_seo import characteristic_rows, normalized, occurrence, plain, SeoError

GREEN, YELLOW, RED, BLUE = 'E2F0D9', 'FFF2CC', 'FCE4D6', '17365D'


def tokens(text):
    return Counter(re.findall(r'\w+', text.casefold().replace('ё', 'е')))


def delta(before, after):
    left, right = tokens(before), tokens(after)
    render = lambda bag: '; '.join(word + (f' ×{n}' if n > 1 else '') for word, n in bag.items()) or 'Нет'
    return render(right - left), render(left - right)


def rich_description(before, after, phrases):
    spans = []
    for phrase in phrases:
        pattern = re.compile(r'(?<!\w)' + re.escape(phrase) + r'(?!\w)', re.I)
        existing = len(list(pattern.finditer(before)))
        spans.extend((m.start(), m.end()) for m in list(pattern.finditer(after))[existing:])
    merged = []
    for start, end in sorted(spans):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(end, merged[-1][1]))
        else:
            merged.append((start, end))
    result, cursor = CellRichText(), 0
    for start, end in merged:
        if start > cursor:
            result.append(after[cursor:start])
        result.append(TextBlock(InlineFont(b=True), after[start:end]))
        cursor = end
    if cursor < len(after):
        result.append(after[cursor:])
    return result if merged else after


def write(cell, value):
    if isinstance(value, (str, CellRichText)) and len(str(value)) > 32767:
        raise SeoError('seo_excel_cell_too_long')
    cell.value = value
    # External/user/model strings are always text, never executable formulas.
    if isinstance(value, str):
        cell.data_type = 's'


def append(sheet, values):
    row = sheet.max_row + 1
    for col, value in enumerate(values, 1):
        cell = sheet.cell(row, col)
        write(cell, value)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    return row


def sheet(wb, name, headers, widths, note):
    ws = wb.create_sheet(name)
    ws.append([name]); ws.append([note]); ws.append([]); ws.append(headers)
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
    ws.cell(1, 1).font = Font(name='Calibri', size=17, bold=True, color=BLUE)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[2].height = 40
    for c in ws[4]:
        c.fill, c.font = PatternFill('solid', fgColor=BLUE), Font(bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[4].height = 32
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'C5'
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = '1:4'
    return ws


def changed(cell, old, new):
    cell.fill = PatternFill('solid', fgColor=GREEN if str(old) == str(new) else YELLOW)
    if str(old) != str(new):
        cell.comment = Comment('Прежнее значение:\n' + str(old), 'OptiCards')


def jump(cell, target, row, label):
    write(cell, label)
    cell.hyperlink = f"#'{target}'!A{row}"
    cell.font = Font(color='0563C1', underline='single')


def build_workbook(path, version, previous=None):
    wb = Workbook(); wb.remove(wb.active)
    note = f"Версия {version['number']}; методика {version['method']}; данные MPStats {version['period']['d1']} — {version['period']['d2']}; свои данные {version.get('sourceCapturedAt', 'см. запуск')} ."
    summary = sheet(wb, 'Итог', ['SKU', 'Артикул', 'Товар', 'Основные проблемы', 'Приоритет', 'Изменений', 'Заголовок', 'Характеристики', 'Описание', 'Вопросов', 'Видимость: начало → конец', 'Дней с наблюдением'], [17,17,40,60,20,12,16,18,16,12,35,24], note)
    sem = sheet(wb, 'Семантика', ['SKU','Запрос','count MPStats','Позиция','Динамика','Релевантность','Куда добавить','Комментарий','Источник','Решение по описанию','У конкурента','В нашем тексте','Реальное вхождение'], [17,38,15,14,30,22,22,65,25,25,65,65,45], note + ' count не объявляется месячной частотой без подтверждения единиц; отсутствие наблюдения — пустая ячейка.')
    titles = sheet(wb, 'Заголовки', ['SKU','Артикул','Текущий заголовок','Знаков','Предлагаемый заголовок','Знаков','Добавленные ключи и слова — полный состав','Удалённые слова — полный состав','Обоснование'], [17,17,60,10,70,10,55,55,65], note)
    chars = sheet(wb, 'Характеристики', ['SKU','Артикул','ID','Характеристика','Текущее значение','Предлагаемое значение','Действие','Обоснование','Источник подтверждения'], [17,17,14,32,55,55,17,65,65], note + ' Фильтр «Действие» сначала показывает правки; сохранённые значения доступны при снятии фильтра. Название/аннотация — на профильных листах.')
    descriptions = sheet(wb, 'Описания', ['SKU','Артикул','Существующее описание','Дополненное описание','Добавленные ключи','Комментарий по изменениям'], [17,17,95,95,45,70], note + ' Новые вхождения ключей выделены настоящим жирным шрифтом.')
    comps = sheet(wb, 'Конкуренты', ['SKU','SKU конкурента','Товар','Дата версии','Почему сопоставим','Исходное описание','Структура','Ключи и решения'], [17,17,55,23,60,95,60,95], note + ' Текст из сохранённой версии MPStats, не гарантированная текущая витрина.')
    issues = sheet(wb, 'Противоречия', ['SKU','Поле','Расхождение','Вопрос / решение','Блокирует принятие'], [17,35,80,80,22], note)
    changes = sheet(wb, 'Изменения версии', ['SKU','Раздел','Поле','Прежнее предложение','Новое предложение','Основание'], [17,22,32,75,75,65], 'Сравнение с предыдущим файлом. Первая версия содержит только исходные предложения на профильных листах.')
    bold_expected, counts = {}, {}
    for card in version['cards']:
        sku, article = str(card['sku']), str(card.get('offerId', ''))
        result = version['results'][sku]
        title, desc = result['title']['value'], result['description']['value']
        before_title, before_desc = card.get('title', ''), plain(card.get('description', ''))
        added, removed = delta(before_title, title)
        title_phrases = [q['query'] for q in result['sources']['semantic'] if occurrence(q['query'], title) and not occurrence(q['query'], before_title)]
        added = 'Поисковые фразы: ' + ('; '.join(title_phrases) or 'Нет нового точного вхождения; см. слова и обоснование') + '\nВсе добавленные слова: ' + added
        tr = append(titles, [sku, article, before_title, len(before_title), title, len(title), added, removed, result['title'].get('reason', '')])
        changed(titles.cell(tr, 5), before_title, title)
        titles.row_dimensions[tr].height = 150
        phrases = list(dict.fromkeys(r['outputPhrase'] for r in result['keywords'] if r['decision'] != 'исключить' and r.get('outputPhrase')))
        rich = rich_description(before_desc, desc, phrases)
        new_phrases = [p for p in phrases if len(re.findall(r'(?<!\w)' + re.escape(p) + r'(?!\w)', desc, re.I)) > len(re.findall(r'(?<!\w)' + re.escape(p) + r'(?!\w)', before_desc, re.I))]
        dr = append(descriptions, [sku, article, before_desc, rich, '; '.join(new_phrases), result['description'].get('reason', '')])
        bold_expected[dr] = rich
        changed(descriptions.cell(dr, 4), before_desc, desc)
        lines = max(sum(max(1, (len(line) + 79) // 80) for line in text.split('\n')) for text in [before_desc, desc])
        descriptions.row_dimensions[dr].height = min(409.5, max(90, lines * 15))
        proposals = {str(r['id']): r for r in result['characteristics']}
        own = characteristic_rows(card)
        existing_ids = {r['id'] for r in own}
        definitions = {str(d['id']): d for d in result['sources']['definitions']}
        all_rows = own + [{'id': aid, 'name': definitions[aid].get('name', aid), 'values': []} for aid in proposals if aid not in existing_ids]
        char_count, first_char = 0, None
        for row in all_rows:
            aid = row['id']
            if card.get('marketplace') != 'Wildberries' and aid in {'4180', '4191'}:
                continue
            proposal = proposals.get(aid)
            old = '; '.join(row['values'])
            new = '; '.join(proposal['values']) if proposal else old
            action = 'оставить' if old == new else ('дополнить' if not old or set(row['values']) < set(proposal['values']) else 'изменить')
            cr = append(chars, [sku, article, aid, row['name'], old, new, action,
                                (proposal.get('reason', '') + ('\nУдаления: ' + '; '.join(f"{r['value']}: {r['reason']}" for r in proposal['removedValues']) if proposal.get('removedValues') else '')) if proposal else 'Корректное исходное значение сохранено.',
                                json_proofs(proposal.get('evidence', [])) if proposal else result['sources'].get('ownSource', 'Свой Ozon Seller API')])
            changed(chars.cell(cr, 6), old, new)
            chars.row_dimensions[cr].height = min(200, max(45, ((max(len(old), len(new)) + 49) // 50) * 15))
            if action == 'оставить':
                chars.row_dimensions[cr].hidden = True
            else:
                char_count += 1
                first_char = first_char or cr
        source = result['sources']
        comp = source['competitor']
        append(comps, [sku, comp['sku'], comp.get('title', ''), comp.get('versionDate', ''), comp.get('reason', '') + '\n' + result.get('compatibility', {}).get('reason', ''), comp['description'], '\n'.join(result['competitorProfile']['structure']),
                       '\n'.join(f"{r['sourcePhrase']} → {r.get('outputPhrase', '')}: {r['decision']}; {r['reason']}" for r in result['keywords'])])
        decisions = {r['query']: r for r in result['semanticDecisions']}
        keyword_map = {r['sourcePhrase']: r for r in result['keywords']}
        competitor_phrases = {normalized(r['phrase']): r for r in result['competitorProfile'].get('phrases', [])}
        queries = {r['query']: r for r in source['semantic']}
        for phrase in keyword_map:
            queries.setdefault(phrase, {'query': phrase, 'source': ''})
        for query, q in queries.items():
            decision, keyword = decisions.get(query, {}), keyword_map.get(query, {})
            origin = competitor_phrases.get(normalized(query))
            origin_text = ''
            if origin:
                origin_text = ('Заголовок' if origin.get('source') == 'title' else 'Описание') + ' конкурента: ' + str(origin.get('quote') or origin['phrase'])
                if origin.get('match') == 'without_parenthetical_synonym':
                    origin_text += ' (сопоставление без вставки в скобках)'
            elif not q.get('source'):
                origin_text = 'Текст конкурента'
            source_text = '; '.join(part for part in [q.get('source', ''), f"count SKU {q['countSourceSku']}" if q.get('countSourceSku') else '', origin_text] if part)
            append(sem, [sku, query, q.get('count'), q.get('position'), q.get('dynamics'), decision.get('relevance', 'По решению из текста конкурента'), decision.get('destination', 'описание' if keyword.get('decision') != 'исключить' else 'не добавлять'), decision.get('reason', keyword.get('reason', '')), source_text, keyword.get('decision', ''), keyword.get('sourcePhrase', ''), keyword.get('outputPhrase', ''), decision.get('witness', '')])
        first_issue = None
        for issue in result['issues']:
            ir = append(issues, [sku, issue.get('field', ''), issue.get('detail', ''), issue.get('question', ''), 'Да' if issue.get('blocking') else 'Нет'])
            first_issue = first_issue or ir
            for c in issues[ir]:
                c.fill = PatternFill('solid', fgColor=RED)
        count = int(before_title != title) + char_count + int(before_desc != desc)
        counts[sku] = count
        visibility = {str(row.get('data', row.get('date', ''))): row for row in source.get('visibility', []) if isinstance(row, dict) and not row.get('no_data')} if isinstance(source.get('visibility', []), list) else {}
        left = visibility.get(version['period']['d1'], {}).get('search_visibility')
        right = visibility.get(version['period']['d2'], {}).get('search_visibility')
        visibility_delta = f"{left if left is not None else 'нет наблюдения'} → {right if right is not None else 'нет наблюдения'}"
        sr = append(summary, [sku, article, before_title, '; '.join(i.get('field', '') for i in result['issues']), 'Есть вопросы' if result['issues'] else 'Проверить', count, '', '', '', len(result['issues']), visibility_delta, len(visibility)])
        summary.cell(sr, 11).comment = Comment('MPStats search_visibility — исходные единицы API, не проценты. no_data не заменяется нулём.', 'OptiCards')
        jump(summary.cell(sr, 7), 'Заголовки', tr, 'Открыть')
        if first_char:
            jump(summary.cell(sr, 8), 'Характеристики', first_char, f'Правок: {char_count}')
        else:
            write(summary.cell(sr, 8), 'Без изменений')
        jump(summary.cell(sr, 9), 'Описания', dr, 'Открыть')
        if first_issue:
            jump(summary.cell(sr, 10), 'Противоречия', first_issue, str(len(result['issues'])))
            summary.cell(sr, 10).fill = PatternFill('solid', fgColor=RED)
        summary.row_dimensions[sr].height = 48
        if previous and sku in previous.get('results', {}):
            old = previous['results'][sku]
            reason = '\n'.join(r['text'] for r in version['remarks'] if r['sku'] in {'*', sku}) or 'Повторная подготовка'
            for section in ('title', 'description'):
                left, right = old[section]['value'], result[section]['value']
                if left != right:
                    rr = append(changes, [sku, section, section, left, right, reason])
                    changed(changes.cell(rr, 5), left, right)
            old_chars = {str(r['id']): '; '.join(r['values']) for r in old['characteristics']}
            new_chars = {str(r['id']): '; '.join(r['values']) for r in result['characteristics']}
            originals = {r['id']: '; '.join(r['values']) for r in own}
            for aid in sorted(old_chars.keys() | new_chars.keys()):
                left, right = old_chars.get(aid, originals.get(aid, '')), new_chars.get(aid, originals.get(aid, ''))
                if left != right:
                    rr = append(changes, [sku, 'characteristics', aid, left, right, reason])
                    changed(changes.cell(rr, 5), left, right)
    for ws in wb:
        if version['cards'] and all(card.get('marketplace') == 'Wildberries' for card in version['cards']):
            ws.cell(4, 1).value = 'nmID'
        ws.auto_filter.ref = f'A4:{get_column_letter(ws.max_column)}{max(4, ws.max_row)}'
        jump(ws.cell(3, ws.max_column), 'Итог', 4, '← Итог')
    chars.auto_filter.add_filter_column(6, ['дополнить', 'изменить'])
    sem.column_dimensions.group('J', 'M', hidden=True)
    summary.column_dimensions['D'].hidden = True
    summary.column_dimensions['D'].outlineLevel = 1
    summary.column_dimensions.group('K', 'L', hidden=True)
    wb.active = 0
    wb.save(path)
    check = load_workbook(path, rich_text=True)
    expected = [str(c['sku']) for c in version['cards']]
    if [str(check['Итог'].cell(r, 1).value) for r in range(5, 5 + len(expected))] != expected or len(set(expected)) != len(expected):
        raise SeoError('seo_export_coverage_mismatch')
    for r, value in bold_expected.items():
        saved = check['Описания'].cell(r, 4).value
        if str(saved) != str(value):
            raise SeoError('seo_export_rich_text_mismatch')
        if isinstance(value, CellRichText):
            import xml.etree.ElementTree as ET
            if not isinstance(saved, CellRichText) or ET.tostring(saved.to_tree()) != ET.tostring(value.to_tree()):
                raise SeoError('seo_export_rich_text_mismatch')
    for ws in check:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    raise SeoError('seo_export_formula_not_allowed')
    if sum(check['Итог'].cell(r, 6).value for r in range(5, 5 + len(expected))) != sum(counts.values()):
        raise SeoError('seo_export_count_mismatch')


def json_proofs(proofs):
    return '\n'.join(f"{p['source']}: {p['quote']}" for p in proofs)
