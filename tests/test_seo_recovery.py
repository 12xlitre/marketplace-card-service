"""Regressions for the 46/58 JEWEL run: bounded work and resumable model stages."""
from collections import Counter
from copy import deepcopy
import json
import threading
import time
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from urllib.error import HTTPError

import server
import automatic_seo
from automatic_seo import AutomaticSeoService, CardCache, SeoError, card_budget_seconds, remember_valid_response, fingerprint
import test_automatic_seo as fixtures


def payload(value, finish='stop'):
    return {'choices': [{'message': {'content': json.dumps(value, ensure_ascii=False)}, 'finish_reason': finish}]}, 'fake', 'fake'


class SeoRecoveryTest(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.AutomaticSeoTest()
        self.case.setUp()
        self.service, self.host = self.case.service, self.case.host

    def tearDown(self):
        self.case.tearDown()

    def test_thousand_queries_split_without_loss_and_reuse_validated_parts(self):
        own = fixtures.card('1001')
        src = fixtures.sources(own, self.case.mapping['1001'])
        src['semantic'] = [{'query': f'запрос {i}'} for i in range(1001)]
        calls = []
        def answer(messages):
            evidence = json.loads(messages[1]['content'])
            queries = evidence.get('semanticBatch', evidence.get('semanticReviewBatch'))
            calls.append(tuple(queries.values()))
            if len(queries) > 6:
                return payload({}, 'length')
            if 'semanticBatch' in evidence:
                return payload({k: {'relevance': 'релевантен', 'reason': 'Назначение подтверждено.'} for k in queries})
            return payload({k: {'destination': 'не добавлять', 'outputPhrase': '',
                               'reason': 'Запрос дублирует уже объяснённое назначение товара в описании.'} for k in queries})
        cache = {}
        progress = []
        with patch.object(self.host, 'audit_llm_chat_completion', side_effect=answer):
            assessments = self.service.assess_semantics(own, src, cache, progress.append)
            result = fixtures.generated()
            result.update(sources=src, semanticDecisions=[dict(r, destination='не добавлять') for r in assessments])
            self.service.explain_unplaced_semantics(own, result, cache, progress.append)
            before = len(calls)
            progress.clear()
            self.assertEqual(assessments, self.service.assess_semantics(own, src, cache, progress.append))
            self.service.explain_unplaced_semantics(own, result, cache, progress.append)
            self.assertEqual(before, len(calls))
            self.assertEqual(progress, [], 'Cached reads must not rewrite the large private snapshot')
        self.assertEqual([r['query'] for r in assessments], [r['query'] for r in src['semantic']])
        self.assertEqual(len(result['semanticDecisions']), 1001)
        self.assertTrue(all(r['placementReviewed'] for r in result['semanticDecisions']))

    def test_failed_additional_plan_resumes_completed_parts_without_paid_repetition(self):
        evidence = {'ownEvidence': {'title': 'Товар'}, 'marketplace': 'Wildberries',
                    'card': {'sku': '1001'}, 'competitor': {}, 'definitions': [],
                    'competitorProfile': {'structure': ['Назначение'], 'phrases': []},
                    'keywordIds': {}, 'semantic': [{'query': f'запрос {i}'} for i in range(101)],
                    'remarks': [], 'planFields': ['title', 'characteristics', 'keywords']}
        completed, fail = Counter(), [True]
        def answer(messages):
            data = json.loads(messages[1]['content'])
            if 'semanticCandidates' not in data:
                return payload({'compatibility': {'comparable': True, 'reason': 'Назначение'},
                                'title': {'value': 'Товар', 'reason': 'Факты'}, 'characteristics': [], 'issues': []})
            queries = tuple(r['query'] for r in data['semanticCandidates'])
            if queries[0] == 'запрос 48' and fail[0]:
                fail[0] = False
                raise RuntimeError('simulated interruption')
            if len(queries) > 6:
                return payload({}, 'length')
            completed[queries] += 1
            return payload({'additionalKeywords': []})
        cache = {}
        with patch.object(self.host, 'audit_llm_chat_completion', side_effect=answer):
            with self.assertRaisesRegex(RuntimeError, 'interruption'):
                self.service.plan_in_parts('План', evidence, {'sku': '1001'}, cache)
            result = self.service.plan_in_parts('План', evidence, {'sku': '1001'}, cache)
        self.assertEqual(result['additionalKeywords'], [])
        self.assertTrue(all(count == 1 for count in completed.values()))
        self.assertEqual({q for batch in completed for q in batch}, {r['query'] for r in evidence['semantic']})

    def test_stage_cache_requires_exact_input_schema_instruction_and_complete_response(self):
        cache = {}
        with patch.object(self.host, 'audit_llm_chat_completion', return_value=payload({'ok': True})) as call:
            for kwargs in [{}, {}, {'evidence': {'changed': True}}, {'instruction': 'changed'},
                           {'response_schema': {'type': 'object'}}, {'raw_text': True}]:
                data = {'instruction': 'test', 'evidence': {}}
                data.update(kwargs)
                self.service.chat(**data, journal=cache, journal_key='stage')
                remember_valid_response(cache, 'stage')
            self.assertEqual(call.call_count, 5)
            cache['stage']['contentLength'] += 1
            self.service.chat('test', {}, journal=cache, journal_key='stage', raw_text=True)
            self.assertEqual(call.call_count, 6)

    def test_syntactically_valid_but_incomplete_answer_is_not_reused_on_retry(self):
        own = fixtures.card('1001')
        src = fixtures.sources(own, self.case.mapping['1001'])
        cache = {}
        with patch.object(self.host, 'audit_llm_chat_completion', side_effect=[payload({}), payload(fixtures.assessment())]) as call:
            with self.assertRaisesRegex(SeoError, 'seo_semantic_decisions_incomplete'):
                self.service.assess_semantics(own, src, cache)
            rows = self.service.assess_semantics(own, src, cache)
        self.assertEqual(call.call_count, 2)
        self.assertEqual(rows[0]['query'], src['semantic'][0]['query'])

    def test_competitor_hallucination_gets_one_repair_with_literal_witness(self):
        competitor = {'sku': '9001', 'title': 'Кран для умывальника', 'description': 'Корпус из латуни.'}
        cache = {}
        with patch.object(self.service, 'chat', side_effect=[
            {'structure': ['Назначение'], 'phrases': [{'phrase': 'выдуманный запрос'}]},
            {'structure': ['Назначение'], 'phrases': [{'phrase': 'Кран для умывальника', 'quote': 'чужая цитата'}]},
        ]) as call:
            self.service.competitor_profile(competitor, cache, 'profile')
        self.assertEqual(call.call_count, 2)
        self.assertEqual(cache['profile']['phrases'][0]['quote'], 'Кран для умывальника')

    def test_official_characteristics_are_checked_before_single_placement_pass(self):
        own = fixtures.card('1001')
        src = fixtures.sources(own, self.case.mapping['1001'])
        events = []
        original = self.service.explain_unplaced_semantics
        def review(*args, **kwargs):
            events.append('placement')
            return original(*args, **kwargs)
        with patch.object(self.service, 'explain_unplaced_semantics', side_effect=review):
            self.service.generate(own, src, [], None, {}, verify_characteristics=lambda result: events.append('official'))
        self.assertEqual(events, ['official', 'placement'])

    def test_worker_does_not_charge_an_hour_waiting_for_model_to_card_budget(self):
        for sku in ('1003', '1004', '1005'):
            self.host.tasks.append({'id': 'task-' + sku, 'cardKey': sku, 'sku': sku, 'offerId': sku, 'status': 'draft'})
            self.case.mapping[sku] = deepcopy(self.case.mapping['1001'])
        elapsed, queued = [0], [0]
        clock = SimpleNamespace(monotonic=lambda: time.monotonic() + elapsed[0])
        self.host.automatic_seo_queue_seconds = lambda: queued[0]
        original = self.service.generate
        def generate(*args, **kwargs):
            if not queued[0]:
                elapsed[0] = queued[0] = 3700
            return original(*args, **kwargs)
        with patch.object(automatic_seo, 'time', clock), patch.object(self.service, 'generate', side_effect=generate):
            self.case.action()
            result = self.case.finish()
        self.assertEqual(result['status'], 'review', result)
        self.assertEqual(result['completedCount'], 5)


class ProviderCooldownTest(unittest.TestCase):
    def test_429_cooldown_is_visible_before_another_worker_enters(self):
        ticks, calls = [100.0], []
        fake_time = SimpleNamespace(monotonic=lambda: ticks[0], sleep=lambda seconds: ticks.__setitem__(0, ticks[0] + seconds))
        def request(*args):
            calls.append(ticks[0])
            if len(calls) == 1:
                raise HTTPError('', 429, 'limit', {'Retry-After': '7'}, None)
            return {}, 'fake', 'fake'
        with patch.object(server, 'time', fake_time), patch.object(server, 'LLM_COOLDOWN_UNTIL', 0), \
             patch.object(server, 'LLM_QUEUE_METRICS', threading.local()), \
             patch.object(server, '_audit_llm_chat_completion', side_effect=request):
            server.audit_llm_chat_completion([])
            self.assertEqual(calls, [100, 107])
            self.assertEqual(server.automatic_seo_queue_seconds(), 7)

    def test_long_retry_after_stops_other_requests_without_contacting_provider(self):
        with patch.object(server, 'LLM_COOLDOWN_UNTIL', 0), \
             patch.object(server, '_audit_llm_chat_completion', side_effect=HTTPError('', 429, 'limit', {'Retry-After': '600'}, None)) as call:
            for _ in range(2):
                with self.assertRaises(HTTPError):
                    server.audit_llm_chat_completion([])
            self.assertEqual(call.call_count, 1)

    def test_budget_scales_for_real_incident_and_has_finite_upper_bound(self):
        self.assertEqual(card_budget_seconds(0), 3600)
        self.assertGreater(card_budget_seconds(961), 3600)
        self.assertEqual(card_budget_seconds(100000), 14400)


class CharacteristicRepairTest(unittest.TestCase):
    setUp = SeoRecoveryTest.setUp
    tearDown = SeoRecoveryTest.tearDown

    def inputs(self):
        own = fixtures.card('1001')
        return own, fixtures.sources(own, self.case.mapping['1001']), {'id':'100','values':[], 'reason':'Неполный ответ', 'evidence':[]}

    def test_keep_original_is_explicit_and_valid_repair_is_cached(self):
        own, src, invalid = self.inputs()
        response = {'action':'оставить исходное','values':[], 'reason':'Материал уже подтверждён собственными данными.', 'evidence':[]}
        cache = {}
        with patch.object(self.host, 'audit_llm_chat_completion', return_value=payload(response)) as call:
            repaired, issue = self.service.repair_characteristic(invalid, own, src, [], cache)
            self.service.repair_characteristic(invalid, own, src, [], cache)
        self.assertIsNone(repaired)
        self.assertIn('латунь', issue['detail'])
        self.assertFalse(issue['blocking'])
        self.assertEqual(call.call_count, 1)
        self.assertEqual(own['characteristics'][1]['values'], ['латунь'])

    def test_decorative_field50_preserves_confirmed_facets(self):
        own, src, _ = self.inputs()
        own['characteristics'].append({'id':50,'name':'Декоративные элементы','values':['Рельефные грани']})
        src['definitions'].append({'id':50,'name':'Декоративные элементы'})
        before = deepcopy(own)
        response = {'action':'оставить исходное','values':[], 'reason':'Рельефные грани подтверждены собственной характеристикой.', 'evidence':[]}
        with patch.object(self.service, 'chat', return_value=response):
            repaired, issue = self.service.repair_characteristic({'id':'50','values':[]}, own, src, [], {})
        self.assertIsNone(repaired)
        self.assertIn('Рельефные грани', issue['detail'])
        self.assertEqual(own, before)

    def test_new_values_require_literal_own_evidence_and_cannot_change_id(self):
        own, src, invalid = self.inputs()
        response = {'id':'foreign-id','action':'предложить значения','values':['латунь'], 'reason':'Собственное описание подтверждает материал.',
                    'evidence':[{'source':'description','quote':'латунь'}]}
        with patch.object(self.service, 'chat', return_value=response) as call:
            repaired, issue = self.service.repair_characteristic(invalid, own, src, [], {})
        self.assertEqual(repaired['id'], '100')
        self.assertEqual(repaired['values'], ['латунь'])
        self.assertEqual(repaired['evidence'][0]['quote'], 'латунь')
        self.assertIsNone(issue)
        self.assertEqual(call.call_count, 1)
        response['values'] = ['пластик']
        with patch.object(self.service, 'chat', return_value=response):
            with self.assertRaisesRegex(SeoError, 'seo_characteristic_evidence_missing'):
                self.service.repair_characteristic(invalid, own, src, [], {})
        response['values'] = ['латунь']
        response['evidence'][0]['quote'] = 'пластик'
        with patch.object(self.service, 'chat', return_value=response):
            with self.assertRaisesRegex(SeoError, 'seo_characteristic_evidence_missing'):
                self.service.repair_characteristic(invalid, own, src, [], {})

    def test_second_malformed_value_stays_an_error_and_offline_recheck_never_pays(self):
        own, src, invalid = self.inputs()
        bad = {'action':'предложить значения','values':[], 'reason':'Пустое значение', 'evidence':[]}
        with patch.object(self.service, 'chat', return_value=bad) as call:
            with self.assertRaisesRegex(SeoError, 'seo_invalid_characteristic_values'):
                self.service.repair_characteristic(invalid, own, src, [], {})
        self.assertEqual(call.call_count, 1)
        with patch.object(self.service, 'chat', side_effect=AssertionError('Offline must not request generation')) as call:
            with self.assertRaisesRegex(SeoError, 'seo_invalid_characteristic_values'):
                self.service.repair_characteristic(invalid, own, src, [], {}, reuse_responses=True)
        call.assert_not_called()

    def test_full_generation_preserves_characteristics_after_format_repair(self):
        original = self.host.audit_llm_chat_completion
        def answer(messages):
            data = json.loads(messages[1]['content'])
            if data.get('repairCharacteristic'):
                return payload({'action':'оставить исходное','values':[],
                                'reason':'Собственный материал уже подтверждён, правка не нужна.','evidence':[]})
            result = original(messages)
            if 'ownEvidence' in data and data.get('card'):
                plan = json.loads(result[0]['choices'][0]['message']['content'])
                plan['characteristics'] = [{'id':'100','values':[], 'reason':'Модель оставила неверный формат','evidence':[]}]
                return payload(plan)
            return result
        with patch.object(self.host, 'audit_llm_chat_completion', side_effect=answer):
            self.case.action()
            finished = self.case.finish()
        self.assertEqual(finished['status'], 'review', finished)
        self.assertEqual(finished['completedCount'], 2)
        version = self.service.load(self.service.directory(1, 'batch'))['versions'][-1]
        for result in version['results'].values():
            self.assertEqual(result['characteristics'], [])
            self.assertTrue(any(i.get('code')=='characteristic_format_repaired' for i in result['issues']))


class CacheAndPollingTest(unittest.TestCase):
    def test_handoff_report_checks_both_original_retry_and_latest_saved_checkpoint(self):
        case = fixtures.AutomaticSeoTest()
        case.setUp()
        try:
            case.host.fail_sku = '1002'
            case.action()
            case.finish()
            case.host.fail_sku = None
            case.action(version=1)
            case.finish()
            directory = case.service.directory(1, 'batch')
            version = case.service.load(directory)['versions'][-1]
            checkpoint = directory / 'checkpoints' / 'handoff-test'
            checkpoint.mkdir(parents=True)
            (checkpoint / 'manifest.json').write_text(json.dumps({'version':2,'done':2,'savedAt':version['finishedAt'],
                'preservedResults':{sku:fingerprint(result) for sku,result in version['results'].items()}}))
            with patch.object(server, 'automatic_seo_service', return_value=case.service), patch.object(server, 'user_is_admin', return_value=True):
                report = server.automatic_seo_handoff_report(1, 'batch', case.tech)
            self.assertEqual(report['unchangedCount'], 1)
            self.assertEqual(report['latestHandoff']['unchangedCount'], 2)
            self.assertTrue(report['latestHandoff']['allPreserved'])
            with patch.object(server, 'user_is_admin', return_value=False):
                with self.assertRaises(PermissionError):
                    server.automatic_seo_handoff_report(1, 'batch', case.tech)
        finally:
            case.tearDown()

    def test_card_cache_copies_only_used_entries_and_isolates_nested_writes(self):
        class Unrelated:
            def __deepcopy__(self, memo):
                raise AssertionError('Another card\'s data must not be copied')
        source = {'other-card': Unrelated(), 'own-card': {'rows': [1]}}
        first, second = CardCache(source), CardCache(source)
        self.assertIn('other-card', first)
        self.assertEqual(len(list(first.materialized_items())), 0)
        first['own-card']['rows'].append(2)
        self.assertEqual(source['own-card']['rows'], [1])
        self.assertEqual(second['own-card']['rows'], [1])
        self.assertEqual(first.baseline_hashes['own-card'], fingerprint({'rows': [1]}))
        source['own-card'] = {'rows': [9]}
        self.assertEqual(first['own-card']['rows'], [1, 2])
        first['new'] = {'rows': []}
        self.assertEqual({key for key, _ in first.materialized_items()}, {'own-card', 'new'})
        del first['other-card']
        self.assertNotIn('other-card', first)
        self.assertEqual(len(first), 2)

    def test_live_metadata_read_does_not_wait_for_private_state_writer(self):
        case = fixtures.AutomaticSeoTest()
        case.setUp()
        release, locked, read = threading.Event(), threading.Event(), threading.Event()
        case.host.gate = threading.Event()
        threads = []
        try:
            case.action()
            directory = case.service.directory(1, 'batch')
            def writer():
                with case.service.lock(directory):
                    locked.set()
                    release.wait(5)
            def reader():
                result = case.service.view(1, 'batch', case.tech)
                self.assertEqual(result['versions'][-1]['status'], 'running')
                read.set()
            threads.append(threading.Thread(target=writer))
            threads[0].start()
            self.assertTrue(locked.wait(2))
            threads.append(threading.Thread(target=reader))
            threads[1].start()
            self.assertTrue(read.wait(1), 'Progress read was blocked by private state persistence')
        finally:
            release.set()
            for thread in threads:
                thread.join(5)
            case.host.gate.set()
            if (case.service.directory(1, 'batch') / 'state.json').exists():
                case.finish()
            case.tearDown()


class LargeBatchRecoveryTest(unittest.TestCase):
    def test_110_card_batch_retries_only_failure_after_service_recreation_and_exports_every_sku(self):
        case = fixtures.AutomaticSeoTest()
        case.setUp()
        try:
            skus = [str(1000 + i) for i in range(110)]
            case.host.tasks = [{'id': 'task-' + sku, 'cardKey': sku, 'sku': sku,
                                'offerId': sku, 'status': 'draft'} for sku in skus]
            case.mapping = {sku: {'sku': '9001', 'reason': 'Сопоставимое назначение.'} for sku in skus}
            case.host.automatic_seo_parallelism = lambda: 4
            case.host.fail_sku = skus[87]
            def finish():
                deadline = time.monotonic() + 120
                while time.monotonic() < deadline:
                    version = case.service.view(1, 'batch', case.tech)['versions'][-1]
                    if version['status'] != 'running':
                        return version
                    time.sleep(.05)
                self.fail('110-card batch exceeded the offline test budget')
            case.action()
            first = finish()
            self.assertEqual(first['status'], 'partial', first)
            self.assertEqual((first['done'], first['total'], first['completedCount']), (110, 110, 109))
            self.assertEqual([e['sku'] for e in first['errors']], [skus[87]])
            directory = case.service.directory(1, 'batch')
            saved = case.service.load(directory)['versions'][-1]
            hashes = {sku: fingerprint(result) for sku, result in saved['results'].items()}
            own_fetches = case.host.own_fetches
            case.host.fail_sku = None
            case.host.calls.clear()
            case.service = AutomaticSeoService(case.host, case.tmp.name)
            resumed = case.action(version=1)['versions'][-1]
            self.assertEqual((resumed['done'], resumed['completedCount']), (109, 109))
            self.assertEqual(resumed['continuation']['preservedCount'], 109)
            self.assertNotIn('_preservation', resumed)
            second = finish()
            self.assertEqual(second['status'], 'review', second)
            self.assertEqual((second['done'], second['total'], second['completedCount']), (110, 110, 110))
            self.assertEqual(case.host.calls, [skus[87]])
            self.assertEqual(case.host.own_fetches, own_fetches)
            final = case.service.load(directory)['versions'][-1]
            self.assertEqual(list(final['results']), skus)
            self.assertEqual(final['sourceCapturedAt'], saved['sourceCapturedAt'])
            self.assertTrue(all(fingerprint(final['results'][sku]) == digest for sku, digest in hashes.items()))
            with patch.object(server, 'automatic_seo_service', return_value=case.service), \
                 patch.object(server, 'user_is_admin', return_value=True):
                report = server.automatic_seo_handoff_report(1, 'batch', case.tech)
            self.assertEqual(report['unchangedCount'], 109)
            self.assertTrue(report['allPreserved'] and report['sourcePreserved'])
            self.assertEqual(report['preservedFromVersion'], 1)
            from io import BytesIO
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(case.service.download(1, 'batch', case.tech, 2)))
            try:
                self.assertEqual(len(book.sheetnames), 8)
                self.assertEqual([str(book['Итог'].cell(i + 5, 1).value) for i in range(110)], skus)
            finally:
                book.close()
        finally:
            case.tearDown()


if __name__ == '__main__':
    unittest.main()
