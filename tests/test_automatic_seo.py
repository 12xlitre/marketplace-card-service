from copy import deepcopy
from datetime import date
import hashlib
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

from automatic_seo import AutomaticSeoService, METHOD, SeoError, closed_period
from automatic_seo_xlsx import build_workbook, delta


def card(sku):
    return {'sku': str(sku), 'offerId': 'TEST-' + str(sku), 'title': 'Смеситель для раковины ТЕСТ',
            'description': 'Смеситель для раковины. Материал: латунь.', 'brand': 'ТЕСТ',
            'characteristics': [{'id': 85, 'name': 'Бренд', 'values': ['ТЕСТ']},
                                {'id': 100, 'name': 'Материал', 'values': ['латунь']}],
            'rawFields': {'attributeDefinitions': [{'id': 85, 'name': 'Бренд'}, {'id': 100, 'name': 'Материал'}]}}


def sources(own, competitor):
    return {'competitor': {'sku': competitor['sku'], 'description': 'Смеситель для раковины — кран для умывальника.',
                           'title': 'Конкурент ТЕСТ', 'versionDate': '2026-10-01', 'reason': competitor['reason']},
            'definitions': own['rawFields']['attributeDefinitions'],
            'semantic': [{'query': 'кран для умывальника', 'count': None, 'position': None, 'source': 'MPStats'}]}


def generated(text='Смеситель для раковины. Латунный кран для умывальника.'):
    return {'compatibility': {'comparable': True, 'reason': 'Сопоставимы назначение и конструкция.'},
            'title': {'value': 'Смеситель для раковины ТЕСТ, латунь', 'reason': 'Добавлен подтверждённый материал.'},
            'draftDescription': text, 'description': {'value': text, 'reason': 'Блоки конкурента, собственный материал.', 'editorChanges': []},
            'characteristics': [], 'keywords': [{'sourcePhrase': 'кран для умывальника', 'outputPhrase': 'кран для умывальника', 'decision': 'использовать', 'reason': 'Соответствует назначению.'}],
            'semanticDecisions': [{'query': 'кран для умывальника', 'relevance': 'релевантен', 'destination': 'описание', 'reason': 'Подтверждено назначение.'}], 'issues': []}


class FakeHost:
    def __init__(self):
        self.tasks = [{'id': 'task-' + sku, 'cardKey': sku, 'sku': sku, 'offerId': sku, 'status': 'draft'} for sku in ['1001', '1002']]
        self.calls, self.own_fetches = [], 0
        self.profile_calls = 0
        self.statuses = []
        self.fail_sku = None
        self.gate = None

    def automatic_seo_context(self, portal, batch, user, edit=False):
        if str(portal) != '1' or batch != 'batch' or user['login'] == 'stranger':
            raise PermissionError('forbidden')
        return {'tasks': deepcopy(self.tasks), 'capabilities': {'prepare': user['login'] == 'tech', 'review': user['login'] == 'manager'}}

    def automatic_seo_own_cards(self, portal, tasks, user):
        self.own_fetches += 1
        if self.gate:
            self.gate.wait(3)
        return [card(t['sku']) for t in tasks]

    def automatic_seo_sources(self, portal, own, competitor, period, cache, user):
        return sources(own, competitor)

    def automatic_seo_verify_characteristics(self, *args):
        pass

    def automatic_seo_set_task_status(self, portal, batch, tasks, status, user):
        self.statuses.append(status)
        for task in self.tasks:
            task['status'] = status

    def audit_llm_chat_completion(self, messages):
        evidence = json.loads(messages[1]['content'])
        if 'ownEvidence' not in evidence:
            self.profile_calls += 1
            answer = {'structure': ['Назначение', 'Материал'], 'phrases': [{'phrase': 'кран для умывальника', 'quote': 'кран для умывальника'}]}
        else:
            sku = evidence['card']['sku']
            self.calls.append(sku)
            if self.fail_sku == sku:
                raise RuntimeError('private-token-must-not-leak')
            text = 'Смеситель для раковины. Латунный кран для умывальника.'
            if evidence['remarks']:
                text += ' Материал — латунь.'
            answer = generated(text)
        return {'choices': [{'message': {'content': json.dumps(answer, ensure_ascii=False)}}]}, 'fake', 'fake'

    def parse_llm_json_content(self, value):
        return json.loads(value)


class AutomaticSeoTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.host = FakeHost()
        self.service = AutomaticSeoService(self.host, self.tmp.name)
        self.tech, self.manager = {'login': 'tech'}, {'login': 'manager'}
        self.mapping = {sku: {'sku': '9001', 'reason': 'Сопоставимы назначение и конструкция.'} for sku in ['1001', '1002']}

    def tearDown(self):
        self.tmp.cleanup()

    def action(self, action='generate', version=0, user=None, **kwargs):
        return self.service.action(1, 'batch', user or self.tech, {'action': action, 'version': version, 'competitors': self.mapping, **kwargs})

    def finish(self):
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            state = self.service.view(1, 'batch', self.tech)
            if state['versions'][-1]['status'] != 'running':
                return state['versions'][-1]
            time.sleep(.01)
        self.fail('Worker did not finish')

    def test_complete_review_return_revision_approval_and_history(self):
        self.action()
        first = self.finish()
        self.assertEqual(first['status'], 'review', first)
        self.assertEqual(self.host.own_fetches, 1)
        self.assertEqual(self.host.profile_calls, 1)
        original = self.service.download(1, 'batch', self.tech, 1)
        self.assertEqual(hashlib.sha256(original).hexdigest(), first['sha256'])
        self.assertEqual(self.action('accept', 1, acknowledgeIssues=True)['versions'][-1]['status'], 'submitted')
        self.action('return', 1, user=self.manager, remarks=[{'sku': '1001', 'section': 'description', 'text': 'Сделать формулировку материала понятнее.'}])
        self.action('generate', 1)
        second = self.finish()
        self.assertEqual(second['status'], 'review', second)
        self.assertEqual(self.host.calls, ['1001', '1002', '1001'])
        self.assertEqual(self.host.own_fetches, 1)
        state = self.service.load(self.service.directory(1, 'batch'))
        old, new = state['versions']
        self.assertEqual(old['results']['1002'], new['results']['1002'])
        self.assertEqual(old['results']['1001']['title'], new['results']['1001']['title'])
        self.assertNotEqual(old['results']['1001']['description'], new['results']['1001']['description'])
        self.assertEqual(self.service.download(1, 'batch', self.tech, 1), original)
        self.action('accept', 2, acknowledgeIssues=True)
        self.action('approve', 2, user=self.manager)
        self.assertEqual(self.host.statuses, ['submitted', 'returned', 'submitted', 'approved'])

    def test_duplicate_running_version_and_access_boundaries(self):
        self.host.gate = threading.Event()
        try:
            self.action()
            with self.assertRaisesRegex(SeoError, 'seo_version_conflict'):
                self.action()
            with self.assertRaisesRegex(SeoError, 'seo_already_running'):
                self.action(version=1)
            with self.assertRaises(PermissionError):
                self.service.download(2, 'batch', self.tech, 1)
            with self.assertRaises(PermissionError):
                self.service.view(1, 'batch', {'login': 'stranger'})
        finally:
            self.host.gate.set()
            self.finish()

    def test_partial_retry_only_failed_and_no_partial_download(self):
        self.host.fail_sku = '1002'
        self.action(); result = self.finish()
        self.assertEqual(result['status'], 'partial')
        self.assertNotIn('private-token', json.dumps(result))
        with self.assertRaises(SeoError):
            self.service.download(1, 'batch', self.tech, 1)
        with self.assertRaisesRegex(SeoError, 'seo_not_ready'):
            self.action('accept', 1)
        self.host.fail_sku = None
        self.action(version=1); self.assertEqual(self.finish()['status'], 'review')
        self.assertEqual(self.host.calls, ['1001', '1002', '1002'])

    def test_stale_accept_roles_and_hash_tampering(self):
        with self.assertRaises(PermissionError):
            self.action(user=self.manager)
        self.action(); self.finish()
        with self.assertRaisesRegex(SeoError, 'seo_review_issues_required'):
            self.action('accept', 1)
        with self.assertRaises(PermissionError):
            self.action('approve', 1)
        with self.assertRaisesRegex(SeoError, 'seo_version_conflict'):
            self.action('accept', 0)
        path = self.service.directory(1, 'batch') / 'seo-v1.xlsx'
        path.write_bytes(b'invalid')
        with self.assertRaisesRegex(SeoError, 'seo_file_verification_failed'):
            self.action('accept', 1, acknowledgeIssues=True)

    def test_scope_change_blocks_download_and_regeneration(self):
        self.action(); self.finish()
        self.host.tasks[0]['cardKey'] = 'changed'
        with self.assertRaisesRegex(SeoError, 'seo_task_scope_changed'):
            self.service.download(1, 'batch', self.tech, 1)
        with self.assertRaisesRegex(SeoError, 'seo_task_scope_changed'):
            self.action(version=1)
        with self.assertRaisesRegex(SeoError, 'seo_task_scope_changed'):
            self.action('accept', version=1, acknowledgeIssues=True)

    def test_explicit_refresh_updates_all_sources_in_new_version(self):
        self.action(); self.finish()
        self.action(version=1, refreshSources=True)
        self.assertEqual(self.finish()['status'], 'review')
        self.assertEqual(self.host.own_fetches, 2)
        self.assertEqual(self.host.calls, ['1001', '1002', '1001', '1002'])
        self.assertEqual(self.host.profile_calls, 2)

    def test_interrupted_run_can_continue_without_repeating_finished_cards(self):
        self.action(); self.finish()
        directory = self.service.directory(1, 'batch')
        state = self.service.load(directory)
        state['versions'][-1]['status'] = 'running'
        self.service.save(directory, state)
        self.assertEqual(self.service.view(1, 'batch', self.tech)['versions'][-1]['status'], 'interrupted')
        self.action(version=1)
        self.assertEqual(self.finish()['status'], 'review')
        self.assertEqual(self.host.calls, ['1001', '1002'])

    def test_lost_key_and_fabricated_characteristic_evidence_rejected(self):
        own, src = card('1001'), sources(card('1001'), self.mapping['1001'])
        result = generated('Латунный смеситель.')
        with patch.object(self.service, 'chat', side_effect=[{'structure': ['Материал'], 'phrases': []}, result]):
            with self.assertRaisesRegex(SeoError, 'seo_keyword_lost_after_editing'):
                self.service.generate(own, src, [], None, {})
        result = generated()
        result['characteristics'] = [{'id': '100', 'values': ['пластик'], 'evidence': [{'source': 'description', 'quote': 'Пластиковый'}]}]
        with patch.object(self.service, 'chat', side_effect=[{'structure': ['Материал'], 'phrases': []}, result]):
            with self.assertRaisesRegex(SeoError, 'seo_characteristic_evidence_missing'):
                self.service.generate(own, src, [], None, {})

    def test_missing_competitor_rejected_before_api_calls(self):
        self.mapping.pop('1002')
        with self.assertRaisesRegex(SeoError, 'seo_competitor_confirmation_required'):
            self.action()
        self.assertEqual(self.host.own_fetches, 0)

    def test_capacity_limit_does_not_start_or_spend_generation(self):
        self.service.slots.acquire(); self.service.slots.acquire()
        try:
            with self.assertRaisesRegex(SeoError, 'seo_capacity_busy'):
                self.action()
            self.assertEqual(self.service.view(1, 'batch', self.tech)['versions'], [])
            self.assertEqual(self.host.own_fetches, 0)
        finally:
            self.service.slots.release(); self.service.slots.release()

    def test_semantic_placement_requires_real_witness(self):
        result = generated()
        result['semanticDecisions'][0]['destination'] = 'характеристика'
        with self.assertRaisesRegex(SeoError, 'seo_semantic_witness_missing'):
            self.service.verify_keywords(result, sources(card('1001'), self.mapping['1001']))

    def test_closed_period_has_30_closed_days(self):
        period = closed_period()
        self.assertEqual((date.fromisoformat(period['d2']) - date.fromisoformat(period['d1'])).days, 29)


class ExportTest(unittest.TestCase):
    def test_large_batches_preserve_all_skus_and_rich_text_and_safe_strings(self):
        with tempfile.TemporaryDirectory() as temp:
            for size in [5, 50, 200]:
                cards = [card(1000+i) for i in range(size)]
                results = {}
                for c in cards:
                    r = generated()
                    r['sources'] = sources(c, {'sku': '9001', 'reason': 'Тестовый пример.'})
                    r['competitorProfile'] = {'structure': ['Назначение']}
                    r['issues'] = [{'field': '=DANGEROUS()', 'detail': 'test', 'question': 'test', 'blocking': False}]
                    results[c['sku']] = r
                version = {'number': 1, 'method': METHOD, 'period': closed_period(), 'cards': cards, 'results': results, 'remarks': []}
                path = Path(temp)/f'test-{size}.xlsx'
                build_workbook(path, version)
                w = load_workbook(path, rich_text=True)
                self.assertEqual(w['Итог'].max_row, size+4)
                self.assertEqual(len(w.sheetnames), 8)
                self.assertEqual(w['Противоречия']['B5'].data_type, 's')
                self.assertIsInstance(w['Описания']['D5'].value, CellRichText)
                self.assertTrue(w['Характеристики'].row_dimensions[5].hidden)
                self.assertEqual(w['Итог']['A5'].value, '1000')
                self.assertEqual(w['Итог'].cell(size+4,1).value, str(999+size))
                self.assertEqual(w['Семантика']['C5'].value, None)

    def test_title_delta_counts_every_added_word_and_repeat(self):
        added, removed = delta('Горшок маленький переносной', 'Горшок маленький переносной, детская накладка на унитаз, горшок')
        self.assertEqual(removed, 'Нет')
        for word in ['детская', 'накладка', 'на', 'унитаз', 'горшок']:
            self.assertIn(word, added)


if __name__ == '__main__':
    unittest.main()
